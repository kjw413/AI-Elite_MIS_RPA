"""Read-only incident audit. Never saves or recalculates input workbooks."""
from pathlib import Path
from datetime import date, datetime
from collections import defaultdict
import json
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DB = Path('E:/DB_MIS')
OUT = ROOT / 'analysis' / 'energy_zero_20260918.json'

def energy(path):
    result = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb:
        for n, row in enumerate(ws.values, 1):
            if row and isinstance(row[0], (date, datetime)):
                result[(ws.title, row[0].strftime('%Y-%m-%d'))] = {'row': n, 'values': list(row[1:19])}
    wb.close()
    return result

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    paths = [DB / 'backup' / name for name in [
        'DB_에너지_backup_20260916_142840.xlsx',
        'DB_에너지_backup_20260917_140047.xlsx',
        'DB_에너지_backup_20260918_154844.xlsx',
    ]] + [DB / 'DB_에너지.xlsx']
    snapshots = [energy(p) for p in paths]
    report = {'files': [str(p) for p in paths], 'comparisons': [], 'current_summary': [], 'raw_headers': []}
    for i in range(1, len(paths)):
        previous, current = snapshots[i-1:i+1]
        changed, zeros = [], []
        for (sheet, day), r in current.items():
            if (sheet, day) not in previous:
                continue
            old = previous[(sheet, day)]
            for j, (a, b) in enumerate(zip(old['values'], r['values']), 2):
                if a != b:
                    rec = {'sheet': sheet, 'date': day, 'column': openpyxl.utils.get_column_letter(j), 'row': r['row'], 'before': a, 'after': b}
                    changed.append(rec)
                    if isinstance(a, (int, float)) and a != 0 and b == 0:
                        zeros.append(rec)
        counts = defaultdict(int)
        for x in zeros:
            counts[x['column']] += 1
        report['comparisons'].append({'before_file': str(paths[i-1]), 'after_file': str(paths[i]), 'changed_count': len(changed), 'zero_transition_columns': dict(counts), 'zero_transitions': zeros, 'all_changes': changed})
        print('COMPARE', paths[i].name, 'changed', len(changed), 'positive_to_zero', dict(counts), flush=True)
    current = snapshots[-1]
    for sheet in dict.fromkeys(k[0] for k in current):
        rows = [(day, r) for (s, day), r in current.items() if s == sheet and '2026-09-01' <= day <= '2026-09-17']
        d = {'sheet': sheet, 'rows': len(rows), 'production_zero_days': [day for day,r in rows if r['values'][12] == 0], 'production_sum': sum(r['values'][12] or 0 for day,r in rows)}
        report['current_summary'].append(d)
        print('CURRENT', d, flush=True)
    rawpaths = [DB / 'backup' / 'RawDB_생산실적_20260917_135055.xlsx', DB / 'backup' / 'RawDB_생산실적_20260918_153836.xlsx', DB / 'RawDB_생산실적.xlsx']
    for p in rawpaths:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for ws in wb:
            rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
            report['raw_headers'].append({'file': str(p), 'sheet': ws.title, 'period': rows[0][:3], 'headers': rows[1]})
        wb.close()
        print('RAW', p.name, report['raw_headers'][-10]['headers'], flush=True)
    raw = energy(DB / 'RawDB_에너지.xlsx')
    report['energy_source_mismatches'] = []
    for key,r in current.items():
        if key not in raw or key[1] < '2026-09-01': continue
        for j,(a,b) in enumerate(zip(raw[key]['values'][:12],r['values'][:12]),2):
            if a != b: report['energy_source_mismatches'].append({'sheet':key[0],'date':key[1],'column':j,'raw':a,'db':b})
    print('ENERGY_SOURCE_MISMATCHES',len(report['energy_source_mismatches']), flush=True)
    wb = openpyxl.load_workbook(DB/'DB_에너지.xlsx',read_only=True,data_only=False)
    report['formula_sample'] = {c: wb['남양주1'][c].value for c in ['N976','O976','Q976','N992','Q992']}
    wb.close()
    OUT.write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print('SAVED',OUT,flush=True)

if __name__ == '__main__':
    main()
