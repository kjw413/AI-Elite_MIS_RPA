"""September missing-range incident: no MIS access or production workbook writes."""
from datetime import date
from unittest.mock import Mock, patch

import pytest
from openpyxl import Workbook

import energy_builder as eb
import production_builder as pb
import production_daily_rpa as rpa
import run_all_rpa
from production_validation import VERIFIED_PERIOD_MARKER, validate_grid_dates


def write_raw(path, days, quantities, *, verified=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "F20_냉동"
    marker = [pb.PERIOD_MARKER, "2026-09-01", "2026-09-17"]
    if verified:
        marker.append(VERIFIED_PERIOD_MARKER)
    ws.append(marker)
    ws.append(["Item Code", "Item 명", "물품대", "누계계획", "누계실적", "누계진척율"]
              + [f"{d:02d}일" for d in days])
    ws.append(["100001", "메로나", 0, 100, sum(quantities), 100] + quantities)
    wb.save(path)
    wb.close()


def write_db(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "daily"
    ws.append(["date", "factory", "item_code", "actual_qty"])
    ws.append([date(2026, 9, 1), "F20", "100001", 406420])
    ws.append([date(2026, 9, 16), "F20", "100001", 300000])
    wb.save(path)
    wb.close()


def test_legacy_last_day_only_raw_preserves_existing_production(tmp_path):
    """Actual incident shape: monthly marker, only 17일 returned by MIS."""
    raw, db = tmp_path / "raw.xlsx", tmp_path / "db.xlsx"
    write_raw(raw, [17], [418157])
    write_db(db)
    merged = eb.merge_production_actuals(db, raw)
    assert merged[("F20", "26-09-01")] == 406420
    assert merged[("F20", "26-09-16")] == 300000
    assert merged[("F20", "26-09-17")] == 418157


def test_verified_full_period_updates_corrections_and_real_zero_days(tmp_path):
    raw, db = tmp_path / "raw.xlsx", tmp_path / "db.xlsx"
    write_raw(raw, [1, 17], [400000, 418157], verified=True)
    write_db(db)
    merged = eb.merge_production_actuals(db, raw)
    assert merged[("F20", "26-09-01")] == 400000
    assert merged[("F20", "26-09-16")] == 0
    assert merged[("F20", "26-09-17")] == 418157
    daily = pb.consolidate_raw_file(raw)
    assert len(daily) == 17
    # The production DW also sees a correction to zero, rather than retaining stale history.
    old = {"F20": daily.loc[daily.date == date(2026, 9, 16),
                           ["date", "item_code", "actual_qty"]].assign(actual_qty=300000)}
    rebuilt = pb._merge_factory_long(old, daily)["F20"]
    assert rebuilt.loc[rebuilt.date == date(2026, 9, 16), "actual_qty"].item() == 0


def test_explicit_zero_is_respected_even_in_legacy_raw(tmp_path):
    raw, db = tmp_path / "raw.xlsx", tmp_path / "db.xlsx"
    write_raw(raw, [1, 16, 17], [0, 300000, 418157])
    write_db(db)
    assert eb.merge_production_actuals(db, raw)[("F20", "26-09-01")] == 0


def test_verified_wrong_grid_period_cannot_write_energy(tmp_path):
    raw, db, out = (tmp_path / n for n in ("raw.xlsx", "db.xlsx", "energy.xlsx"))
    write_raw(raw, [18], [100], verified=True)
    write_db(db)
    original = b"existing energy file must remain untouched"
    out.write_bytes(original)
    with pytest.raises(ValueError, match="그리드 일자 오류"):
        eb.write_raw({"김해": {date(2026, 9, 17): {"total_power_kwh": 123}}},
                     out, production_path=db, production_raw_path=raw, recalculate=False)
    assert out.read_bytes() == original


@pytest.mark.parametrize("end,expected", [
    # September 3, 2018 was Monday. Include the previous Friday (August 31).
    (date(2018, 9, 2), [(date(2018, 8, 1), date(2018, 8, 31)),
                        (date(2018, 9, 1), date(2018, 9, 2))]),
    (date(2026, 1, 2), [(date(2025, 12, 1), date(2025, 12, 31)),
                        (date(2026, 1, 1), date(2026, 1, 2))]),
    (date(2024, 3, 2), [(date(2024, 2, 1), date(2024, 2, 29)),
                        (date(2024, 3, 1), date(2024, 3, 2))]),
    (date(2026, 8, 31), [(date(2026, 8, 1), date(2026, 8, 31))]),
    (date(2026, 9, 17), [(date(2026, 9, 1), date(2026, 9, 17))]),
])
def test_month_boundary_and_midmonth_ranges(end, expected):
    assert rpa.plan_collection_periods(end, {}) == expected


def test_previous_month_gap_is_recovered_after_overlap_window():
    assert rpa.plan_collection_periods(date(2026, 9, 17), {
        "F20_냉동": {date(2026, 8, 28), date(2026, 9, 16)},
    }) == [(date(2026, 8, 1), date(2026, 8, 31)), (date(2026, 9, 1), date(2026, 9, 17))]


def make_rpa():
    obj = object.__new__(rpa.MISProductionRPA)
    obj.start_date, obj.end_date = "2026-09-01", "2026-09-17"
    obj.coords = {"start_date_field": [609, 108], "end_date_field": [609, 130]}
    obj.main_window = Mock()
    return obj


def test_wrong_start_field_retries_then_stops_before_verified_marker():
    obj = make_rpa()
    with patch.object(obj, "_write_date_field") as write, patch.object(
        obj, "_read_date_field", return_value=date(2026, 9, 17)
    ):
        with pytest.raises(ValueError, match="조회 기간 불일치"):
            obj.set_date_range()
    assert write.call_count == 6
    assert not obj._date_range_verified


def test_explicit_field_clicks_replace_tab_and_verified_dates_allow_collection():
    obj = make_rpa()
    with patch.object(rpa, "fast_click") as click, patch.object(rpa, "send_keys") as keys, \
         patch.object(rpa.time, "sleep"), patch.object(
             obj, "_read_date_field", side_effect=[date(2026, 9, 1), date(2026, 9, 17)]):
        obj.set_date_range()
    assert obj._date_range_verified
    assert {tuple(c.args[1:]) for c in click.call_args_list} == {(609, 108), (609, 130)}
    assert not any(c.args == ("{TAB}",) for c in keys.call_args_list)


def test_date_reader_rejects_stale_clipboard():
    obj = make_rpa()
    with patch.object(rpa, "fast_click"), patch.object(rpa, "send_keys"), \
         patch.object(rpa.time, "sleep"), patch.object(rpa, "get_clipboard_sequence", return_value=7), \
         patch.object(rpa, "wait_for_clipboard_change", return_value=False), \
         patch.object(rpa, "get_clipboard_text", return_value="2026-09-01") as read:
        with pytest.raises(ValueError, match="날짜 필드 읽기 실패"):
            obj._read_date_field("start_date_field")
    read.assert_not_called()


def test_partial_collection_does_not_proceed_to_dw():
    obj = rpa.MISProductionRPA(date_from="2026-09-01", date_to="2026-09-17")
    obj.main_window = Mock()
    target = {"sheet_name": "F20_냉동", "factory": "F20", "category": "냉동"}
    with patch.object(rpa, "discover_targets", return_value=[target]), \
         patch.object(rpa.time, "sleep"), patch.object(obj, "connect_mis"), \
         patch.object(obj, "navigate_to_production_screen"), patch.object(obj, "backup_raw_file"), \
         patch.object(obj, "set_date_range"), patch.object(obj, "select_item_type"), \
         patch.object(obj, "_collect_period", return_value=(1, 1, 5)), \
         patch.object(obj, "_capture_raw_snapshot") as capture:
        assert obj.collect() is False
        assert obj.process_collected_data() is False
    capture.assert_not_called()


def test_cli_keeps_blank_dates_automatic_and_returns_failure_status():
    with patch.object(rpa, "_setup_logging"), patch.object(rpa.sys, "argv", ["production_daily_rpa.py"]), \
         patch.object(rpa, "MISProductionRPA") as ctor:
        ctor.return_value.run.return_value = False
        assert rpa.main() == 1
    assert ctor.call_args.kwargs["date_from"] is None
    assert ctor.call_args.kwargs["date_to"] is None


def test_verified_grid_requires_production_header():
    with pytest.raises(ValueError, match="Item Code"):
        validate_grid_dates(["일자", "01일"], date(2026, 9, 1), date(2026, 9, 17))


def test_full_run_propagates_automatic_month_recovery_to_all_pipelines():
    import sys
    import types
    from argparse import Namespace
    created = {}

    def fake_class(kind):
        class Pipeline:
            def __init__(self, **kwargs):
                created[kind] = kwargs
                self.app, self.main_window = object(), Mock()
                self.collection_periods = []
            def collect(self):
                if kind == "production":
                    self.collection_periods = [(date(2018, 8, 1), date(2018, 8, 31)),
                                               (date(2018, 9, 1), date(2018, 9, 2))]
                return True
            def attach_existing_window(self, app, main_window):
                pass
            def process_collected_data(self):
                return True
        return Pipeline

    modules = {
        "production_daily_rpa": types.SimpleNamespace(MISProductionRPA=fake_class("production")),
        "utility_daily_rpa": types.SimpleNamespace(MISUtilityRPA=fake_class("utility")),
        "wip_daily_rpa": types.SimpleNamespace(MISWIPRPA=fake_class("wip")),
    }
    args = Namespace(date=None, date_from=None, date_to=None, factories=None, dry_run=False)
    with patch.dict(sys.modules, modules), patch.object(
        run_all_rpa.argparse.ArgumentParser, "parse_known_args", return_value=(args, [])
    ):
        assert run_all_rpa.main() == 0
    assert created["production"]["date_from"] is None
    assert created["production"]["date_to"] is None
    for name in ["utility", "wip"]:
        assert created[name]["date_from"] == "2018-08-01"
        assert created[name]["date_to"] == "2018-09-02"


def test_verified_raw_does_not_invent_zero_for_other_namyangju_factory(tmp_path):
    raw = tmp_path / "raw.xlsx"
    write_raw(raw, [1, 17], [100, 200], verified=True)
    from openpyxl import load_workbook
    wb = load_workbook(raw)
    wb.active.title = "F10_냉동"
    wb.save(raw)
    wb.close()
    actuals = eb._load_raw_production_actuals(raw)
    assert actuals[("F10B", "26-09-01")] == 100
    assert actuals[("F10B", "26-09-16")] == 0
    assert not any(factory == "F10A" for factory, _ in actuals)
